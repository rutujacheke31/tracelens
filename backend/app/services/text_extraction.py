from pathlib import Path

from docx import Document
from pypdf import PdfReader

from app.services.storage import UPLOAD_ROOT


PLAIN_TEXT_EXTENSIONS = {
    ".log",
    ".txt",
    ".csv",
    ".json",
}

def extract_plain_text(file_path: Path) -> str:
    """Extract text from LOG, TXT, CSV, and JSON files."""
    return file_path.read_text(
        encoding="utf-8",
        errors="replace",
    )


def extract_pdf_text(file_path: Path) -> str:
    """Extract text from every readable page of a PDF."""
    reader = PdfReader(str(file_path))
    page_texts = []

    for page in reader.pages:
        page_text = page.extract_text()

        if page_text:
            page_texts.append(page_text)

    return "\n\n".join(page_texts)


def extract_docx_text(file_path: Path) -> str:
    """Extract paragraphs and table contents from a DOCX file."""
    document = Document(str(file_path))
    text_parts = []

    for paragraph in document.paragraphs:
        paragraph_text = paragraph.text.strip()

        if paragraph_text:
            text_parts.append(paragraph_text)

    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            )

            if row_text:
                text_parts.append(row_text)

    return "\n".join(text_parts)


def extract_text_from_upload(storage_key: str) -> str:
    """Locate a stored upload and use the correct extraction method."""
    upload_root = UPLOAD_ROOT.resolve()
    file_path = (upload_root / storage_key).resolve()

    # Prevent a storage key from accessing files outside uploads/.
    try:
        file_path.relative_to(upload_root)
    except ValueError as error:
        raise ValueError("Invalid upload storage key.") from error

    if not file_path.is_file():
        raise FileNotFoundError("The uploaded file could not be found.")

    extension = file_path.suffix.lower()

    if extension in PLAIN_TEXT_EXTENSIONS:
        extracted_text = extract_plain_text(file_path)

    elif extension == ".pdf":
        extracted_text = extract_pdf_text(file_path)

    elif extension == ".docx":
        extracted_text = extract_docx_text(file_path)

    elif extension == ".doc":
        raise ValueError(
            "Legacy .doc extraction is not supported yet. "
            "Please convert the file to .docx."
        )

    else:
        raise ValueError(
            f"Text extraction is not supported for {extension} files."
        )

    # Return all extracted text. Chunking is handled separately.
    return extracted_text.strip()